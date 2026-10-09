"""Exports Excel/CSV pour les fichiers de contrôle utilisateur.

Ce module ne contient aucune écriture Odoo. Il met en forme les DataFrames issus
du workflow en onglets de revue lisibles, avec les formats numériques attendus.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from lmc_po_price.models import WorkflowResult
from lmc_po_price.workflow import ambiguous_review, blocked_price_changes, unmatched_review


def result_workbook_bytes(
    result: WorkflowResult,
    include_purchase_order: bool = True,
    include_price_review: bool = True,
    include_developer_sheets: bool = False,
    price_dry_run_results: pd.DataFrame | None = None,
    po_dry_run_results: pd.DataFrame | None = None,
    po_creation_results: pd.DataFrame | None = None,
) -> bytes:
    """Construit le classeur de contrôle selon les tâches sélectionnées."""
    sheets = {
        "tous_les_articles_facture": result.all_lines,
        "articles_trouves": result.matched,
        "articles_non_retrouves": unmatched_review(result.unmatched),
        "articles_ambigus": ambiguous_review(result.ambiguous),
    }
    if include_price_review:
        sheets["changements_prix"] = result.price_changes
        blocked = blocked_price_changes(result.price_changes)
        if not blocked.empty:
            sheets["changements_prix_bloques"] = blocked
        if price_dry_run_results is not None and not price_dry_run_results.empty:
            sheets["dry_run_prix_odoo"] = price_dry_run_results
    if include_purchase_order:
        sheets["bon_commande_a_verifier"] = result.purchase_order_review
        if po_dry_run_results is not None and not po_dry_run_results.empty:
            sheets["dry_run_po_odoo"] = po_dry_run_results
        if po_creation_results is not None and not po_creation_results.empty:
            sheets["creation_po_odoo"] = po_creation_results
        if include_developer_sheets:
            numeric_po = _purchase_order_with_numeric_product_ids(result)
            if not numeric_po.empty:
                sheets["bon_commande_id_numerique"] = numeric_po
    sheets["notes"] = _notes()
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=name[:31])
        _format_workbook(writer)
    return buffer.getvalue()


def _notes() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"champ": "Fact_PU_Net_GZ", "description": "Prix facture augmente du taux de contribution transport, par exemple 0,6%."},
            {"champ": "Fact_PU_unitaire", "description": "Prix facture converti dans l'unite Odoo lorsque la base fournit un ratio d'unite fournisseur."},
            {"champ": "statut", "description": "trouve, non_trouve ou a_verifier."},
            {"champ": "peut_etre_mis_a_jour", "description": "Vrai si le match est certain et si le prix facture differe du prix Odoo a la deuxieme decimale."},
            {"champ": "changements_prix_bloques", "description": "Changements de prix dont le prix de vente ne peut pas etre calcule automatiquement; lire raison_revue."},
            {"champ": "ecart_anormal", "description": "Signalement visuel quand l'ecart depasse le seuil choisi; ne bloque pas la mise a jour Odoo."},
            {"champ": "Odoo_sale_ok", "description": "Valeur a envoyer vers Odoo pour sale_ok lors d'une future mise a jour prix eligible."},
            {"champ": "db_vente_ok", "description": "Champ Odoo sale_ok renomme en francais dans la base locale."},
            {"champ": "dry_run_prix_odoo", "description": "Rapport de verification Odoo sans ecriture avant mise a jour des prix."},
            {"champ": "dry_run_po_odoo", "description": "Rapport de verification Odoo sans creation avant creation du bon de commande."},
            {"champ": "creation_po_odoo", "description": "Rapport de creation reelle du bon de commande dans Odoo."},
            {"champ": "bon_commande_id_numerique", "description": "Onglet diagnostic: meme bon de commande avec Article/ID remplace par l'ID numerique product.product."},
        ]
    )


def _purchase_order_with_numeric_product_ids(result: WorkflowResult) -> pd.DataFrame:
    if result.purchase_order_review.empty or result.matched.empty:
        return pd.DataFrame()

    output = result.purchase_order_review.reset_index(drop=True).copy()
    matched = result.matched.reset_index(drop=True)
    row_count = min(len(output), len(matched))
    output = output.iloc[:row_count].copy()
    matched = matched.iloc[:row_count]
    output.insert(
        output.columns.get_loc("Lignes de la commande/Article/ID") + 1,
        "Lignes de la commande/Article/ID externe origine",
        output["Lignes de la commande/Article/ID"],
    )
    output["Lignes de la commande/Article/ID"] = matched["db_article_id"].map(_format_odoo_numeric_id)
    return output


def _format_odoo_numeric_id(value: object) -> str:
    if pd.isna(value):
        return ""
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return str(value)


def _format_workbook(writer: pd.ExcelWriter) -> None:
    number_formats = {
        "Fact_PU_Net": "0.00",
        "Fact_PU_Net_GZ": "0.00",
        "Fact_PU_unitaire": "0.00",
        "DB_Prix_Net": "0.00",
        "Ecart_Prix": "0.00",
        "Ecart_Prix_percent": "0.0%",
        "New_Prix_de_vente": "0.00",
        "Fact_montant_HT": "0.00",
        "db_cout": "0.00",
        "db_prix_fournisseur": "0.00",
    }
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        headers = [cell.value for cell in ws[1]]
        for column_cells in ws.columns:
            header = str(column_cells[0].value or "")
            width = max(12, min(48, len(header) + 2))
            ws.column_dimensions[column_cells[0].column_letter].width = width
            if header in number_formats:
                for cell in column_cells[1:]:
                    cell.number_format = number_formats[header]
            if header == "Matchs_possibles":
                ws.column_dimensions[column_cells[0].column_letter].width = 80
                for cell in column_cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        _highlight_abnormal_price_rows(ws, headers)


def _highlight_abnormal_price_rows(ws, headers: list[object]) -> None:
    if "ecart_anormal" not in headers:
        return
    abnormal_col = headers.index("ecart_anormal") + 1
    fill = PatternFill(fill_type="solid", fgColor="F8D7DA")
    font = Font(color="842029")
    for row in ws.iter_rows(min_row=2):
        if row[abnormal_col - 1].value is True:
            for cell in row:
                cell.fill = fill
                cell.font = font
